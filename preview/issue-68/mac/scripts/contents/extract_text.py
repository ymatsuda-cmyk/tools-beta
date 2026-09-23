#!/usr/bin/env python3
"""文書ファイルから本文テキストを取り出す。

対応: PDF / テキスト / Word(.docx) / Excel(.xlsx, .xlsm) / PowerPoint(.pptx)

PDFはまず埋め込まれたテキスト層を読む。スキャンしただけのPDFは
テキスト層が無く空になるので、そのときだけOCRに回す(遅いため最後の手段)。

必要なライブラリ (足りないものはその形式だけスキップされる):
    pip install pypdf python-docx openpyxl python-pptx
OCRを使う場合 (PDFを画像にする poppler はどちらでも必要):
    brew install poppler
    macOS なら swift が入っていれば標準のVisionを使う(追加導入不要)
    それ以外は  brew install tesseract tesseract-lang

単体で試す:
    python3 extract_text.py ファイルパス
"""
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".log", ".json", ".yml", ".yaml"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx"}

# OCRに回すかどうかの目安。1ページあたりこれ未満ならテキスト層が無いとみなす
OCR_MIN_CHARS_PER_PAGE = 20
OCR_LANG = os.environ.get("OCR_LANG", "jpn+eng")
OCR_ENGINE = os.environ.get("OCR_ENGINE", "auto")  # auto | vision | tesseract | off
OCR_MAX_PAGES = int(os.environ.get("OCR_MAX_PAGES", "50"))


def is_supported(path):
    return Path(path).suffix.lower() in SUPPORTED_SUFFIXES


def squeeze(text):
    """空行の連続と行末の空白を詰める。要約に渡す前提なので体裁だけ整える。"""
    lines = [re.sub(r"[ \t]+$", "", line) for line in str(text or "").splitlines()]
    out = []
    for line in lines:
        if not line.strip() and out and not out[-1].strip():
            continue
        out.append(line)
    return "\n".join(out).strip()


# ---------------------------------------------------------------- 形式ごとの読み出し


def read_plain(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8", "cp932", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_docx(path):
    from docx import Document

    doc = Document(str(path))
    out = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return "\n".join(out)


def read_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n## シート: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)


def read_pptx(path):
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n## スライド {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        out.append(" | ".join(cells))
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if notes.strip():
            out.append(f"(ノート) {notes.strip()}")
    return "\n".join(out)


def read_pdf_text(path):
    """テキスト層だけを読む。pypdf が無いときは空で返し、OCR側で補完する。"""
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", 0, "pypdf"

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001  1ページ壊れていても残りは読む
            pages.append("")
    return "\n".join(pages), len(reader.pages), None


# ---------------------------------------------------------------- OCR


def ocr_available():
    """使えるOCRの名前。macOSのVisionを優先する(追加導入が要らず日本語に強い)"""
    if OCR_ENGINE == "off":
        return None
    if OCR_ENGINE in ("vision", "auto") and shutil.which("swift"):
        return "vision"
    if OCR_ENGINE in ("tesseract", "auto") and shutil.which("tesseract"):
        return "tesseract"
    return None


def pdf_to_images(path, dest_dir):
    """pdftoppm(poppler)でページを画像にする。無ければ空リスト。"""
    if not shutil.which("pdftoppm"):
        return []
    prefix = str(Path(dest_dir) / "page")
    cmd = ["pdftoppm", "-png", "-r", "200", "-l", str(OCR_MAX_PAGES), str(path), prefix]
    try:
        subprocess.run(cmd, capture_output=True, timeout=900, check=True)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
        return []
    return sorted(Path(dest_dir).glob("page*.png"))


VISION_SWIFT = r'''
import Foundation
import Vision
import AppKit

guard CommandLine.arguments.count > 1,
      let image = NSImage(contentsOfFile: CommandLine.arguments[1]),
      let cg = image.cgImage(forProposedRect: nil, context: nil, hints: nil) else { exit(1) }
let request = VNRecognizeTextRequest()
request.recognitionLevel = .accurate
request.recognitionLanguages = ["ja-JP", "en-US"]
request.usesLanguageCorrection = true
try? VNImageRequestHandler(cgImage: cg, options: [:]).perform([request])
guard let results = request.results as? [VNRecognizedTextObservation] else { exit(0) }
for observation in results {
    if let line = observation.topCandidates(1).first { print(line.string) }
}
'''


def build_vision_tool(dest_dir, log):
    """Visionを叩く小さな実行ファイルを1回だけ作る。

    swift で都度解釈させるとページごとに数秒かかるため、先にコンパイルする。
    swiftc が無ければ swift でその都度実行する。
    """
    source = Path(dest_dir) / "ocr.swift"
    source.write_text(VISION_SWIFT)
    if not shutil.which("swiftc"):
        return ["swift", str(source)]
    binary = Path(dest_dir) / "ocr"
    try:
        subprocess.run(["swiftc", "-O", "-o", str(binary), str(source)],
                       capture_output=True, timeout=300, check=True)
        return [str(binary)]
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        log(f"    ⚠️ Vision用のビルドに失敗したため swift で実行します: {type(e).__name__}")
        return ["swift", str(source)]


def ocr_image(image_path, engine, vision_cmd=None):
    cmd = (["tesseract", str(image_path), "stdout", "-l", OCR_LANG]
           if engine == "tesseract" else [*vision_cmd, str(image_path)])
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        return res.stdout if res.returncode == 0 else ""
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return ""


def ocr_pdf(path, log=print):
    engine = ocr_available()
    if not engine:
        log("    ⚠️ OCRの準備ができていません(brew install tesseract tesseract-lang poppler)")
        return ""
    with tempfile.TemporaryDirectory() as tmp:
        images = pdf_to_images(path, tmp)
        if not images:
            log("    ⚠️ PDFを画像にできませんでした(brew install poppler)")
            return ""
        vision_cmd = build_vision_tool(tmp, log) if engine == "vision" else None
        log(f"    → OCR中({engine}) {len(images)}ページ...")
        return "\n".join(ocr_image(p, engine, vision_cmd) for p in images)


# ---------------------------------------------------------------- 入口


def extract(path, log=print):
    """@returns (text, engine_label)。読めなければ (None, 理由)"""
    path = Path(path)
    suffix = path.suffix.lower()

    try:
        if suffix in TEXT_SUFFIXES:
            return squeeze(read_plain(path)), "テキスト"
        if suffix == ".docx":
            return squeeze(read_docx(path)), "Word"
        if suffix in (".xlsx", ".xlsm"):
            return squeeze(read_xlsx(path)), "Excel"
        if suffix == ".pptx":
            return squeeze(read_pptx(path)), "PowerPoint"
        if suffix == ".pdf":
            text, pages, missing_lib = read_pdf_text(path)
            text = squeeze(text)
            if pages and len(text) >= pages * OCR_MIN_CHARS_PER_PAGE:
                return text, "PDF"
            log("    テキスト層がほぼ空 → OCRにフォールバック")
            ocr = squeeze(ocr_pdf(path, log))
            if ocr and (len(ocr) > len(text) or missing_lib == "pypdf"):
                return ocr, "PDF(OCR)"
            if missing_lib == "pypdf" and not ocr:
                return None, "ライブラリが未導入です (pypdf)"
            return text, "PDF"
    except ImportError as e:
        return None, f"ライブラリが未導入です ({e.name})"
    except Exception as e:  # noqa: BLE001  壊れたファイルで止めない
        return None, f"{type(e).__name__}: {e}"

    return None, f"未対応の形式です ({suffix or '拡張子なし'})"


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("使い方: python3 extract_text.py ファイルパス", file=sys.stderr)
        sys.exit(1)

    target = Path(sys.argv[1]).expanduser()
    print(f"OCRエンジン: {ocr_available() or '使えません'}")
    body, label = extract(target)
    if body is None:
        print(f"❌ {label}", file=sys.stderr)
        sys.exit(1)
    print(f"✅ {label} / {len(body):,}字\n{'-' * 60}")
    print(body[:2000])
    if len(body) > 2000:
        print(f"\n... (残り {len(body) - 2000:,}字)")
